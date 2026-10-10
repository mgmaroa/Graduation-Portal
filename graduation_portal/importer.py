"""
Reading the registrar's graduation list (CSV or Excel) into the database.

Three steps, kept separate so the ICT chief can preview before anything changes:
    parse_upload()  -> read the file, validate every row
    build_plan()    -> work out what would be added / updated / left alone
    apply_import()  -> do it, all-or-nothing, and write the UploadLog

Uploads only ADD and UPDATE. They never remove real students, because later
lists contain students who were cleared afterwards. Records edited by hand are
protected from being overwritten unless the ICT chief chooses otherwise.
"""
import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime

from django.conf import settings
from django.db import transaction
from django.utils import timezone

from .models import Graduand, UploadLog
from .utils import normalize_admission_no

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 20000
CHUNK = 500  # keeps "IN (...)" queries under SQLite/MySQL parameter limits

REQUIRED = ("admission_no", "student_name", "course", "school")
LABELS = {
    "student_name": "Name",
    "course": "Course",
    "school": "School / department",
    "graduation_date": "Graduation date",
}
MAX_LEN = {"student_name": 150, "course": 200, "school": 200}

# Header text (lowercased, letters and digits only) -> our field
HEADER_ALIASES = {
    "admission_no": {"admissionno", "admissionnumber", "admno", "admnumber", "admission",
                     "studentid", "studentno", "studentnumber"},
    "student_name": {"studentname", "name", "fullname", "names", "nameofstudent", "namesofstudent"},
    "course": {"course", "courseofstudy", "coursename", "programme", "program"},
    "school": {"school", "department", "schooldepartment", "schoolordepartment",
               "schoolname", "departmentname", "faculty"},
    "graduation_date": {"graduationdate", "dateofgraduation", "graduation", "date"},
}
DATE_FORMATS = ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d %B %Y", "%d %b %Y")


class UploadError(Exception):
    """Problem with the file as a whole (wrong type, missing columns...)."""


@dataclass
class ParseResult:
    rows: list = field(default_factory=list)       # valid rows
    rejected: list = field(default_factory=list)   # (row number, reason)
    total: int = 0                                  # data rows seen (valid + rejected)


# ---------------------------------------------------------------------------
# 1. Parse
# ---------------------------------------------------------------------------
def _read_table(uploaded_file):
    name = (uploaded_file.name or "").lower()
    if uploaded_file.size > MAX_FILE_BYTES:
        raise UploadError("The file is larger than 5 MB.")
    raw = uploaded_file.read()

    if name.endswith(".csv"):
        text = None
        for encoding in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return list(csv.reader(io.StringIO(text), dialect))

    if name.endswith(".xlsx"):
        from openpyxl import load_workbook
        try:
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            return [list(r) for r in workbook.active.iter_rows(values_only=True)]
        except Exception:
            raise UploadError("That Excel file could not be read. Re-save it as .xlsx and try again.")

    raise UploadError("Upload a .csv or .xlsx file. For an old .xls file, open it in Excel and use Save As .xlsx.")


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return " ".join(str(value).split())


def _parse_date(value, default):
    if value is None or _text(value) == "":
        return default
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _map_headers(header_row):
    mapping = {}
    for index, cell in enumerate(header_row):
        key = re.sub(r"[^a-z0-9]", "", _text(cell).lower())
        for field_name, aliases in HEADER_ALIASES.items():
            if key in aliases and field_name not in mapping:
                mapping[field_name] = index
    missing = [f for f in REQUIRED if f not in mapping]
    if missing:
        found = ", ".join(_text(c) for c in header_row if _text(c)) or "none"
        names = {"admission_no": "admission number", "student_name": "student name",
                 "course": "course", "school": "school or department"}
        raise UploadError(
            "Missing column(s): " + ", ".join(names[m] for m in missing)
            + f". Columns found in the file: {found}."
        )
    return mapping


def parse_upload(uploaded_file) -> ParseResult:
    table = _read_table(uploaded_file)
    start = next((i for i, row in enumerate(table) if any(_text(c) for c in row)), None)
    if start is None:
        raise UploadError("The file is empty.")

    mapping = _map_headers(table[start])
    default_date = date.fromisoformat(settings.GRADUATION_DATE)
    result = ParseResult()
    seen = {}

    def cell(row, name):
        index = mapping.get(name)
        return row[index] if index is not None and index < len(row) else None

    for index in range(start + 1, len(table)):
        row, number = table[index], index + 1
        if not any(_text(c) for c in row):
            continue
        result.total += 1
        if result.total > MAX_ROWS:
            raise UploadError(f"The file has more than {MAX_ROWS} rows. Split it into smaller files.")

        adm = normalize_admission_no(_text(cell(row, "admission_no")))
        name, course, school = (_text(cell(row, f)) for f in ("student_name", "course", "school"))
        grad_date = _parse_date(cell(row, "graduation_date"), default_date)

        reason = None
        if not re.fullmatch(settings.ADMISSION_NO_PATTERN, adm):
            reason = f"Invalid admission number '{adm or 'blank'}' (use 4 to 6 digits)"
        elif not name:
            reason = "Student name is blank"
        elif not course:
            reason = "Course is blank"
        elif not school:
            reason = "School or department is blank"
        elif grad_date is None:
            reason = f"Unreadable graduation date '{_text(cell(row, 'graduation_date'))}'"
        elif any(len(v) > MAX_LEN[k] for k, v in
                 (("student_name", name), ("course", course), ("school", school))):
            reason = "A value is too long"
        elif adm in seen:
            reason = f"Duplicate admission number (first seen on row {seen[adm]})"

        if reason:
            result.rejected.append((number, reason))
            continue
        seen[adm] = number
        result.rows.append({
            "row": number, "admission_no": adm, "student_name": name, "course": course,
            "school": school, "graduation_date": grad_date.isoformat(),
        })
    return result


# ---------------------------------------------------------------------------
# 2. Plan
# ---------------------------------------------------------------------------
@dataclass
class Plan:
    new: list = field(default_factory=list)
    update: list = field(default_factory=list)          # {"row", "obj", "changes"}
    skipped_manual: list = field(default_factory=list)  # {"row", "obj", "changes"}
    unchanged: int = 0
    dummy_clashes: int = 0      # demo records sharing a number with a real one
    dummy_total: int = 0        # demo records currently in the system


def _existing_map(numbers):
    found = {}
    for i in range(0, len(numbers), CHUNK):
        for g in Graduand.objects.filter(admission_no__in=numbers[i:i + CHUNK]):
            found[g.admission_no] = g
    return found


def _changes(obj, row):
    out = []
    for name, label in LABELS.items():
        old, new = getattr(obj, name), row[name]
        old = old.isoformat() if isinstance(old, date) else old
        if old != new:
            out.append((label, old, new))
    return out


def build_plan(rows, *, overwrite_manual=False, remove_dummy=True) -> Plan:
    plan = Plan(dummy_total=Graduand.objects.filter(is_dummy=True).count())
    existing = _existing_map([r["admission_no"] for r in rows])

    for row in rows:
        obj = existing.get(row["admission_no"])
        if obj is not None and obj.is_dummy:
            if remove_dummy:
                plan.dummy_clashes += 1
                obj = None  # the demo record is deleted first, so this is a fresh record
        if obj is None:
            plan.new.append(row)
            continue

        changes = _changes(obj, row)
        item = {"row": row, "obj": obj, "changes": changes}
        if obj.is_dummy:
            plan.update.append(item)           # keeping demo rows: real data takes over
        elif not changes:
            plan.unchanged += 1
        elif obj.manually_edited and not overwrite_manual:
            plan.skipped_manual.append(item)
        else:
            plan.update.append(item)
    return plan


# ---------------------------------------------------------------------------
# 3. Apply
# ---------------------------------------------------------------------------
def apply_import(rows, rejected, total, *, filename, user, overwrite_manual=False, remove_dummy=True):
    now = timezone.now()
    with transaction.atomic():
        dummy_removed = 0
        if remove_dummy and rows:
            dummy_removed = Graduand.objects.filter(is_dummy=True).delete()[0]

        plan = build_plan(rows, overwrite_manual=overwrite_manual, remove_dummy=False)

        Graduand.objects.bulk_create([
            Graduand(
                admission_no=r["admission_no"], student_name=r["student_name"],
                course=r["course"], school=r["school"],
                graduation_date=date.fromisoformat(r["graduation_date"]),
            ) for r in plan.new
        ], batch_size=CHUNK)

        to_update = []
        for item in plan.update:
            obj, r = item["obj"], item["row"]
            obj.student_name, obj.course, obj.school = r["student_name"], r["course"], r["school"]
            obj.graduation_date = date.fromisoformat(r["graduation_date"])
            obj.is_dummy = False
            if overwrite_manual:
                obj.manually_edited = False
            obj.updated_at = now
            to_update.append(obj)
        Graduand.objects.bulk_update(
            to_update,
            ["student_name", "course", "school", "graduation_date",
             "is_dummy", "manually_edited", "updated_at"],
            batch_size=CHUNK,
        )

        notes = [f"Row {n}: {why}" for n, why in rejected[:200]]
        if len(rejected) > 200:
            notes.append(f"...and {len(rejected) - 200} more rejected rows.")
        if plan.skipped_manual:
            notes.append("Left alone (edited by hand): "
                         + ", ".join(i["row"]["admission_no"] for i in plan.skipped_manual[:200]))
        if plan.unchanged:
            notes.append(f"{plan.unchanged} record(s) already matched and were unchanged.")

        log = UploadLog.objects.create(
            uploaded_by=user, filename=filename[:255], records_total=total,
            records_added=len(plan.new), records_updated=len(plan.update),
            records_skipped=len(plan.skipped_manual), records_rejected=len(rejected),
            dummy_removed=dummy_removed, notes="\n".join(notes),
        )
    return log